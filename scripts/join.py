"""Produce auditable hypothesis tables, never infer convergence from constraint."""
from __future__ import annotations

from collections import defaultdict
import csv
import gzip
import hashlib
import json
import re

import openpyxl
from fetch import DATA, OUT, sha
from fetch_genebass import MASKS, PHENOTYPES, load_json

SNP_THRESHOLD = 5e-8  # Lee 2018 Table 13; Savage 2018 Table S5
MAGMA_THRESHOLD = 2.76e-6  # Savage Table S15, 18,128 tests
SKATO_THRESHOLD = 2.5e-7  # Karczewski 2022 empirical per-phenotype threshold
BURDEN_THRESHOLD = 6.7e-7  # Same paper, distinct burden-test threshold


def tsv(name, rows, fields=None):
    rows = list(rows)
    fields = fields or list(rows[0])
    path = OUT / name
    with path.open("w") as f:
        w = csv.DictWriter(f, fieldnames=fields, delimiter="\t", lineterminator="\n")
        w.writeheader()
        w.writerows({k: ("NA" if row.get(k) is None or row.get(k) == "" else row[k]) for k in fields} for row in rows)
    if path.stat().st_size >= 5_000_000:
        raise RuntimeError(f"Derived table exceeds 5 MB: {path}")


def read_tsv(name):
    return csv.DictReader((DATA / name).open(), delimiter="\t")


def main():
    # Verify exact downloaded bytes before using them; ignore retained error bodies.
    inventory = json.loads((OUT / "fetch-manifest.json").read_text())
    required = {"lee2018.txt", "savage2018.txt", "savage-checksum.txt", "savage-supplement.xlsx", "lee-supplement.xlsx", "gencode-v19.gtf.gz", "hgnc.tsv", "unknome-18_Mar_2026.tsv.gz", "pharos400.csv", "gene2pubmed.gz", "gnomad-v4.1.tsv", "genebass-phenotypes.json.gz", "genebass-phenotype-qc.json.gz"}
    required.update(f"genebass-{p}-{m}.json.gz" for p in PHENOTYPES for m in MASKS)
    required.update(f"genebass-gene-qc-{m}.json.gz" for m in MASKS)
    observed = {r["file"]: r for r in inventory}
    for name in required:
        r = observed.get(name, {})
        if r.get("http_status") != 200 or r.get("expected_sha256"):
            raise RuntimeError(f"Required input unavailable or changed: {name}")
    for r in inventory:
        if r.get("http_status") == 200:
            if sha(DATA / r["file"]) != r["sha256"]:
                raise RuntimeError(f"Input checksum mismatch: {r['file']}")
    md5 = hashlib.md5()  # publisher supplies MD5; inventory additionally has SHA256
    with (DATA / "savage2018.txt").open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            md5.update(chunk)
    if md5.hexdigest() not in (DATA / "savage-checksum.txt").read_text():
        raise RuntimeError("Savage publisher checksum mismatch")

    hgnc = [r for r in read_tsv("hgnc.tsv") if r["status"] == "Approved"]
    mapping_issues = []
    def unique_map(key):
        grouped = defaultdict(list)
        for r in hgnc:
            if r[key]:
                grouped[r[key]].append(r)
        for identifier, matches in grouped.items():
            if len(matches) > 1:
                mapping_issues.append({"source": "HGNC " + key, "id": identifier, "symbol": ";".join(r["symbol"] for r in matches), "reason": "Ambiguous approved-gene identifier; not assigned"})
        return {k: v[0] for k, v in grouped.items() if len(v) == 1}
    ens = unique_map("ensembl_gene_id")
    entrez = unique_map("entrez_id")
    uniprot = defaultdict(set)
    for r in hgnc:
        if r["ensembl_gene_id"] in ens:
            for acc in r["uniprot_ids"].split("|"):
                if acc:
                    uniprot[acc].add(r["ensembl_gene_id"])
    common = defaultdict(dict)
    w = openpyxl.load_workbook(DATA / "savage-supplement.xlsx", read_only=True, data_only=True)
    low_loci = {int(r[0]) for r in w["Table S7"].iter_rows(min_row=6, values_only=True) if isinstance(r[0], (float, int))}
    low_intervals = [(str(int(r[1])), int(r[2]), int(r[3])) for r in w["Table S5"].iter_rows(min_row=7, values_only=True) if isinstance(r[0], (float, int)) and int(r[0]) in low_loci]
    for r in w["Table S12"].iter_rows(min_row=6, values_only=True):
        if not r[2]:
            continue
        gid = str(r[2]).split(".")[0]
        loci = {int(x) for x in str(r[0]).split(":")}
        if loci <= low_loci:
            continue
        common[gid].update(savage_fuma=True, savage_fuma_min_snp_p=r[17], savage_fuma_loci=str(r[0]), savage_fuma_mapping=";".join(m for m, yes in [("positional", bool(r[8])), ("eQTL", bool(r[10])), ("chromatin", r[15] == "Yes")] if yes))
    for r in w["Table S15"].iter_rows(min_row=6, values_only=True):
        if not isinstance(r[0], (float, int)) or not isinstance(r[6], (float, int)):
            continue
        eid = str(int(r[0]))
        if eid not in entrez or not entrez[eid]["ensembl_gene_id"]:
            mapping_issues.append({"source": "Savage S15", "id": eid, "symbol": r[7], "reason": "No current HGNC Entrez-to-Ensembl match"})
            continue
        if r[6] < MAGMA_THRESHOLD:
            common[entrez[eid]["ensembl_gene_id"]]["savage_magma_p"] = r[6]

    # EA3 published gene scores are a separate PROXY-only analysis, never treated
    # as cognitive-performance MAGMA. Table 29 N~1M confirms the distinction.
    ea = {}
    lee = openpyxl.load_workbook(DATA / "lee-supplement.xlsx", read_only=True, data_only=True)
    for r in lee["29. MAGMA Genes"].iter_rows(min_row=3, values_only=True):
        if not isinstance(r[0], (float, int)):
            continue
        h = entrez.get(str(int(r[0])))
        if h and h["ensembl_gene_id"] and isinstance(r[12], (float, int)) and r[12] < 0.05:
            ea[h["ensembl_gene_id"]] = {"ea3_magma_joint_p": r[9], "ea3_magma_fdr": r[12]}

    meta = {r["analysis_id"]: r for r in load_json(DATA / "genebass-phenotypes.json.gz")}
    pheno_qc = {r["analysis_id"]: r for r in load_json(DATA / "genebass-phenotype-qc.json.gz")}
    associations = defaultdict(list)
    audit = []
    coverage = []
    for mask, annotation in MASKS.items():
        gene_qc = {r["gene_id"]: r for r in load_json(DATA / f"genebass-gene-qc-{mask}.json.gz")}
        for p in PHENOTYPES:
            raw = load_json(DATA / f"genebass-{p}-{mask}.json.gz")
            coverage.append({"analysis_id": p, "description": meta[p]["description"], "n": meta[p]["n_cases"], "mask": annotation, "rows": len(raw), **pheno_qc[p]})
            for r in raw:
                gid = r["gene_id"]
                q = gene_qc.get(gid, {})
                ac = q["CAF"] * meta[p]["n_cases"] if q.get("CAF") is not None else None
                # Use paper lower-bound lambda rule; browser flags also impose
                # an upper bound 1.5. Store flags separately rather than silently
                # describing that additional browser threshold as the paper's.
                base = q.get("keep_gene_coverage") is True and q.get("keep_gene_n_var") is True and ac is not None and ac >= 50
                ok = {}
                for test in ["skato", "burden"]:
                    lam = q.get(f"synonymous_lambda_gc_{test}")
                    ok[test] = bool(base and lam is not None and lam >= 0.75 and pheno_qc[p].get(f"keep_pheno_{test}") is True)
                skato = r.get("Pvalue")
                burden = r.get("Pvalue_Burden")
                hit = bool((skato is not None and skato < SKATO_THRESHOLD and ok["skato"]) or (burden is not None and burden < BURDEN_THRESHOLD and ok["burden"]))
                record = {"ensembl_gene_id": gid, "source_symbol": r["gene_symbol"], "analysis_id": p, "rare_mask": annotation, "rare_skato_p": skato, "rare_burden_p": burden, "rare_burden_beta": r.get("BETA_Burden"), "expected_ac": ac, "coverage_pass": q.get("keep_gene_coverage"), "n_variants_pass": q.get("keep_gene_n_var"), "synonymous_lambda_skato": q.get("synonymous_lambda_gc_skato"), "synonymous_lambda_burden": q.get("synonymous_lambda_gc_burden"), "browser_gene_skato_pass": q.get("keep_gene_skato"), "browser_gene_burden_pass": q.get("keep_gene_burden"), "rare_skato_qc_pass": ok["skato"], "rare_burden_qc_pass": ok["burden"], "rare_significant_qc_pass": hit}
                associations[gid].append(record)
                if (skato is not None and skato < SKATO_THRESHOLD) or (burden is not None and burden < BURDEN_THRESHOLD):
                    audit.append(record)
    rare_hits = {gid for gid, rows in associations.items() if any(r["rare_significant_qc_pass"] for r in rows)}

    # Body-only GRCh37 mapping for BOTH direct-cognition GWAS. No nearest-gene
    # heuristic and no arbitrary flanking window. Overlapping bodies all count.
    genes_by_chr = defaultdict(list)
    for line in gzip.open(DATA / "gencode-v19.gtf.gz", "rt"):
        if line.startswith("#"):
            continue
        cols = line.rstrip().split("\t")
        if cols[2] != "gene":
            continue
        gid = re.search(r'gene_id "([^"]+)"', cols[8]).group(1).split(".")[0]
        genes_by_chr[cols[0].removeprefix("chr")].append((int(cols[3]), int(cols[4]), gid))
    snp_counts = {}
    for name, pc, snpc in [("lee2018", "Pval", "MarkerName"), ("savage2018", "P", "SNP")]:
        significant = defaultdict(list)
        counts = {"total": 0, "gws": 0, "gws_body_mapped": 0}
        for r in read_tsv(name + ".txt"):
            counts["total"] += 1
            p = float(r[pc])
            if p < SNP_THRESHOLD:
                counts["gws"] += 1
                pos = int(r["POS"])
                if name != "savage2018" or not any(c == r["CHR"] and start <= pos <= end for c, start, end in low_intervals):
                    significant[r["CHR"]].append((pos, p, r[snpc]))
        for chrom, genes in genes_by_chr.items():
            snps = sorted(significant[chrom])
            active = []
            ordered = sorted(genes)
            idx = 0
            for pos, p, snp in snps:
                while idx < len(ordered) and ordered[idx][0] <= pos:
                    active.append(ordered[idx]); idx += 1
                active = [g for g in active if g[1] >= pos]
                if active:
                    counts["gws_body_mapped"] += 1
                for start, end, gid in active:
                    key = name + "_body_min_snp_p"
                    if p < common[gid].get(key, 1):
                        common[gid].update({key: p, name + "_body_snp": snp})
        snp_counts[name] = counts

    universe = set(common) | rare_hits | {r["ensembl_gene_id"] for r in audit}
    for gid in sorted(universe - ens.keys()):
        mapping_issues.append({"source": "Evidence universe", "id": gid, "symbol": "", "reason": "No unique current HGNC Ensembl mapping; Entrez annotations unavailable"})
    # Unknome score already represents the best-known orthologue in a cluster.
    # HGNC UniProt links join human proteins; max across proteins is conservative.
    scores = defaultdict(list)
    with gzip.open(DATA / "unknome-18_Mar_2026.tsv.gz", "rt") as f:
        for r in csv.DictReader(f, delimiter="\t"):
            if r["taxon_id"] != "9606":
                continue
            for acc in re.split(r"[;,|\s]+", r["uniprot_accessions"]):
                if len(uniprot.get(acc, set())) == 1:
                    gid = next(iter(uniprot[acc]))
                    scores[gid].append((float(r["knownness"]), r["cluster_id"]))
    knownness = {gid: max(x[0] for x in values) for gid, values in scores.items()}
    sorted_scores = sorted(knownness.values())
    ranks = {}
    for i, score in enumerate(sorted_scores, 1):
        ranks.setdefault(score, i)  # competition ranks: equal scores share rank
    pharos = defaultdict(list)
    with (DATA / "pharos400.csv").open() as f:
        for r in csv.DictReader(f):
            # The misleading ensembl_id column contains ENSP protein IDs.
            h = entrez.get(r["ncbi_id"])
            if h and h["ensembl_gene_id"]:
                pharos[h["ensembl_gene_id"]].append(r)
    counts = defaultdict(set)
    needed_entrez = {r["entrez_id"] for r in ens.values() if r["entrez_id"]}
    with gzip.open(DATA / "gene2pubmed.gz", "rt") as f:
        for r in csv.DictReader(f, delimiter="\t"):
            if r["#tax_id"] == "9606" and r["GeneID"] in needed_entrez:
                counts[r["GeneID"]].add(r["PubMed_ID"])
    publication_counts = {gid: len(counts[h["entrez_id"]]) for gid, h in ens.items() if h["entrez_id"]}
    publication_ranks = {}
    for i, count in enumerate(sorted(publication_counts.values()), 1):
        publication_ranks.setdefault(count, i)
    constraint = defaultdict(list)
    for r in read_tsv("gnomad-v4.1.tsv"):
        if r["gene_id"].startswith("ENSG") and r["canonical"] == "true":
            constraint[r["gene_id"]].append(r)

    def annotate(gid):
        h = ens.get(gid, {})
        pr = pharos.get(gid, [])
        cr = constraint.get(gid, [])
        return {"ensembl_gene_id": gid, "symbol": h.get("symbol", ""), "entrez_id": h.get("entrez_id", ""), "unknome_knownness": knownness.get(gid, ""), "unknome_rank_least_known_first": ranks.get(knownness.get(gid), ""), "unknome_rank_universe_n": len(knownness), "unknome_clusters": ";".join(sorted({x[1] for x in scores.get(gid, [])})), "pharos_tdl": ";".join(sorted({r["tdl"] for r in pr})), "pharos_uniprot": ";".join(sorted({r["uniprot_id"] for r in pr})), "human_pubmed_count": publication_counts.get(gid, ""), "pubmed_rank_least_published_first": publication_ranks.get(publication_counts.get(gid), ""), "pubmed_rank_universe_n": len(publication_counts), "gnomad_canonical_transcripts": ";".join(sorted(r["transcript"] for r in cr)), "gnomad_lof_oe_ci_upper": ";".join(sorted({r["lof.oe_ci.upper"] for r in cr})), "gnomad_constraint_flags": ";".join(sorted({r["constraint_flags"] for r in cr}))}

    common_fields = ["savage_magma_p", "savage_fuma", "savage_fuma_min_snp_p", "savage_fuma_loci", "savage_fuma_mapping", "lee2018_body_min_snp_p", "lee2018_body_snp", "savage2018_body_min_snp_p", "savage2018_body_snp"]
    rows = []
    for gid in sorted(universe):
        base = annotate(gid)
        base.update({k: common.get(gid, {}).get(k, "") for k in common_fields})
        base.update(common_direct_significant=gid in common, rare_significant_qc_pass=gid in rare_hits, hypothesis_only=True)
        eligible = [r for r in associations.get(gid, []) if r["rare_significant_qc_pass"]]
        choices = eligible or associations.get(gid, [])
        best = min(choices, key=lambda r: min(r["rare_skato_p"] if r["rare_skato_p"] is not None else 1, r["rare_burden_p"] if r["rare_burden_p"] is not None else 1)) if choices else {}
        for k in ["analysis_id", "rare_mask", "rare_skato_p", "rare_burden_p", "rare_burden_beta", "expected_ac", "rare_skato_qc_pass", "rare_burden_qc_pass"]:
            base[k] = best.get(k, "")
        rows.append(base)
    tsv("gene-evidence.tsv", rows)
    tsv("cognition-hypotheses.tsv", [r for r in rows if r["common_direct_significant"] and r["rare_significant_qc_pass"]], list(rows[0]))
    tsv("rare-threshold-audit.tsv", [{**annotate(r["ensembl_gene_id"]), **r} for r in audit])
    tsv("genebass-coverage.tsv", coverage)
    proxy = [{**annotate(gid), **ea[gid], "common_phenotype": "educational attainment (EA3), NOT direct cognition", "rare_phenotype": meta[r["analysis_id"]]["description"], **r, "hypothesis_only": True} for gid in sorted(rare_hits & ea.keys()) for r in associations[gid] if r["rare_significant_qc_pass"]]
    if proxy:
        tsv("ea-proxy-hypotheses.tsv", proxy)
    tsv("mapping-issues.tsv", mapping_issues, ["source", "id", "symbol", "reason"])
    tsv("knownness-ranks.tsv", [{"ensembl_gene_id": g, "symbol": ens.get(g, {}).get("symbol", ""), "knownness": knownness[g], "rank_least_known_first": ranks[knownness[g]]} for g in sorted(knownness)])
    summary = {"source_thresholds": {"snp": SNP_THRESHOLD, "savage_magma": MAGMA_THRESHOLD, "genebass_skato": SKATO_THRESHOLD, "genebass_burden": BURDEN_THRESHOLD, "expected_ac_min": 50, "synonymous_lambda_min": 0.75, "ea3_proxy_fdr": 0.05}, "snp_counts": snp_counts, "savage_low_confidence_loci_excluded_from_fuma": sorted(low_loci), "common_direct_gene_count": len(common), "rare_significant_qc_gene_count": len(rare_hits), "rare_significant_qc_genes": sorted(rare_hits), "direct_convergence_gene_count": len(set(common) & rare_hits), "ea_proxy_convergence_genes": sorted(rare_hits & ea.keys()), "unknome_rank_universe_n": len(knownness), "gene_evidence_rows": len(rows), "mapping_issues": len(mapping_issues), "inventoried_download_bytes": sum(r.get("bytes", 0) for r in inventory)}
    (OUT / "analysis-summary.json").write_text(json.dumps(summary, indent=2) + "\n")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
